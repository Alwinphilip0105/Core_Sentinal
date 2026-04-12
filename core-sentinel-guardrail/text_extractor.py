"""
Extract text from common document and image types (PDF, Office, plain text, OCR).
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

SUPPORTED = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tiff",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
}


def _empty_result(
    *,
    error: str | None,
    method: str = "unsupported",
    path: Path | None = None,
) -> dict:
    out: dict = {
        "text": "",
        "pages": [],
        "method": method,
        "file_name": path.name if path else "",
        "file_size": path.stat().st_size if path and path.is_file() else 0,
        "page_count": 0,
        "error": error,
    }
    return out


def extract(file_path: str) -> dict:
    """
    Returns:
    {
      "text": full extracted text,
      "pages": list of {"page": N, "text": "..."},
      "method": how it was extracted,
      "file_name": basename,
      "file_size": bytes,
      "page_count": N,
      "error": None or error message
    }
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext not in SUPPORTED:
        return _empty_result(error=f"Unsupported: {ext}", path=path if path.exists() else None)

    if not path.is_file():
        return _empty_result(error=f"Not a file or missing: {file_path}", path=path)

    try:
        if ext == ".pdf":
            return _extract_pdf(path)
        if ext == ".docx":
            return _extract_docx(path)
        if ext in (".xlsx", ".xls"):
            return _extract_excel(path)
        if ext in (".pptx", ".ppt"):
            return _extract_pptx(path)
        if ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            return _extract_image(path)
        return _extract_text(path)
    except Exception as e:
        return {
            "error": str(e),
            "text": "",
            "pages": [],
            "method": "failed",
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": 0,
        }


def _extract_pdf(path: Path) -> dict:
    import pdfplumber

    pages = []
    full_text = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            pages.append({"page": i, "text": t})
            full_text.append(t)

    text = "\n".join(full_text)
    method = "pdf_text"

    if len(text.strip()) < 50:
        return _extract_pdf_ocr(path)

    return {
        "text": text,
        "pages": pages,
        "method": method,
        "page_count": len(pages),
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _extract_pdf_ocr(path: Path) -> dict:
    try:
        from pdf2image import convert_from_path
        import pytesseract

        _set_tesseract()
        images = convert_from_path(str(path), dpi=200)
        pages = []
        for i, img in enumerate(images, 1):
            t = pytesseract.image_to_string(img)
            pages.append({"page": i, "text": t})
        text = "\n".join(p["text"] for p in pages)
        return {
            "text": text,
            "pages": pages,
            "method": "pdf_ocr",
            "page_count": len(pages),
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "error": None,
        }
    except Exception as e:
        return {
            "error": str(e),
            "text": "",
            "pages": [],
            "method": "pdf_ocr_failed",
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": 0,
        }


def _extract_docx(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    texts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            texts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_texts: list[str] = []
            for cell in row.cells:
                if cell.text.strip():
                    row_texts.append(cell.text.strip())
            if row_texts:
                texts.append(" | ".join(row_texts))
    text = "\n".join(texts)
    pages = [{"page": 1, "text": text}]
    return {
        "text": text,
        "pages": pages,
        "method": "docx",
        "page_count": 1,
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _extract_excel(path: Path) -> dict:
    import openpyxl

    ext = path.suffix.lower()
    if ext == ".xls":
        return {
            "error": "Legacy .xls is not supported by openpyxl; convert to .xlsx or install xlrd/pandas.",
            "text": "",
            "pages": [],
            "method": "excel_unsupported",
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": 0,
        }

    wb = openpyxl.load_workbook(str(path), data_only=True)
    pages = []
    all_text = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        sheet_text = "\n".join(rows)
        pages.append({"page": sheet.title, "text": sheet_text})
        all_text.append(sheet_text)
    text = "\n".join(all_text)
    return {
        "text": text,
        "pages": pages,
        "method": "excel",
        "page_count": len(wb.worksheets),
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _extract_pptx(path: Path) -> dict:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    ext = path.suffix.lower()
    if ext == ".ppt":
        return {
            "error": "Legacy .ppt (binary) is not supported; use .pptx or convert the file.",
            "text": "",
            "pages": [],
            "method": "ppt_unsupported",
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": 0,
        }

    try:
        prs = Presentation(str(path))
    except PackageNotFoundError as e:
        return {
            "error": str(e),
            "text": "",
            "pages": [],
            "method": "pptx_failed",
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": 0,
        }

    pages = []
    all_text = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        slide_text = "\n".join(texts)
        pages.append({"page": i, "text": slide_text})
        all_text.append(slide_text)
    text = "\n".join(all_text)
    return {
        "text": text,
        "pages": pages,
        "method": "pptx",
        "page_count": len(prs.slides),
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _extract_image(path: Path) -> dict:
    import pytesseract
    from PIL import Image

    _set_tesseract()
    img = Image.open(str(path))
    text = pytesseract.image_to_string(img)
    pages = [{"page": 1, "text": text}]
    return {
        "text": text,
        "pages": pages,
        "method": "ocr",
        "page_count": 1,
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _extract_text(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    pages = [{"page": 1, "text": text}]
    return {
        "text": text,
        "pages": pages,
        "method": "plaintext",
        "page_count": 1,
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "error": None,
    }


def _set_tesseract() -> None:
    import pytesseract

    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for c in candidates:
        if os.path.exists(c):
            pytesseract.pytesseract.tesseract_cmd = c
            return


if __name__ == "__main__":
    if len(sys.argv) > 1:
        result = extract(sys.argv[1])
        print("Method:", result.get("method"))
        print("Pages:", result.get("page_count"))
        print("Error:", result.get("error"))
        print("Text preview:", result.get("text", "")[:200])
    else:
        print("Usage: python text_extractor.py <file>")
